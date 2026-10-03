import docx
from pathlib import Path
from parsers.base import BaseParser, ParsedDocument, DocumentPage
from core.hash_utils import compute_file_hash
from core.logger import setup_logger

logger = setup_logger("DocxParser")

class DocxParser(BaseParser):
    def can_parse(self, file_path: Path) -> bool:
        return file_path.suffix.lower() == ".docx" and not file_path.name.startswith("~$")

    def parse(self, file_path: Path) -> ParsedDocument:
        path = Path(file_path)
        stat = path.stat()
        file_hash = compute_file_hash(path)

        paragraphs = []
        if stat.st_size == 0:
            return ParsedDocument(
                file_path=str(path.resolve()),
                file_name=path.name,
                extension=path.suffix.lower(),
                file_size=0,
                last_modified=stat.st_mtime,
                content_hash=file_hash,
                full_text="",
                pages=[DocumentPage(page_number=1, text="")],
                is_scanned=False,
                metadata={"empty": True}
            )

        try:
            doc = docx.Document(str(path))
            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    paragraphs.append(text)

            # Also extract table text if present
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
        except Exception as e:
            logger.debug(f"Could not read docx {path.name}: {e}")
            return ParsedDocument(
                file_path=str(path.resolve()),
                file_name=path.name,
                extension=path.suffix.lower(),
                file_size=stat.st_size,
                last_modified=stat.st_mtime,
                content_hash=file_hash,
                full_text=f"[Unreadable Word Document: {e}]",
                pages=[DocumentPage(page_number=1, text=f"[Unreadable: {e}]")],
                is_scanned=False,
                metadata={"error": str(e)}
            )

        full_text = "\n\n".join(paragraphs)

        return ParsedDocument(
            file_path=str(path.resolve()),
            file_name=path.name,
            extension=path.suffix.lower(),
            file_size=stat.st_size,
            last_modified=stat.st_mtime,
            content_hash=file_hash,
            full_text=full_text,
            pages=[DocumentPage(page_number=1, text=full_text)],
            is_scanned=False,
            metadata={"paragraph_count": len(paragraphs)}
        )
