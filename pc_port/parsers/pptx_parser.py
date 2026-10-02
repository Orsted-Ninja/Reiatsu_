import pptx
from pathlib import Path
from parsers.base import BaseParser, ParsedDocument, DocumentPage
from core.hash_utils import compute_file_hash
from core.logger import setup_logger

logger = setup_logger("PptxParser")

class PptxParser(BaseParser):
    def can_parse(self, file_path: Path) -> bool:
        return file_path.suffix.lower() == ".pptx" and not file_path.name.startswith("~$")

    def parse(self, file_path: Path) -> ParsedDocument:
        path = Path(file_path)
        stat = path.stat()
        file_hash = compute_file_hash(path)

        pages = []
        full_text_parts = []
        slide_count = 0

        if stat.st_size == 0:
            return ParsedDocument(
                file_path=str(path.resolve()),
                file_name=path.name,
                extension=path.suffix.lower(),
                file_size=0,
                last_modified=stat.st_mtime,
                content_hash=file_hash,
                full_text="",
                pages=[DocumentPage(page_number=1, text="")],
                is_scanned=False,
                metadata={"empty": True}
            )

        try:
            prs = pptx.Presentation(str(path))
            slide_count = len(prs.slides)
            for idx, slide in enumerate(prs.slides):
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        clean = shape.text.strip()
                        if clean:
                            slide_texts.append(clean)

                # Slide notes if available
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        slide_texts.append(f"[Notes: {notes}]")

                slide_content = "\n".join(slide_texts)
                pages.append(DocumentPage(page_number=idx + 1, text=slide_content))
                if slide_content:
                    full_text_parts.append(f"--- Slide {idx + 1} ---\n{slide_content}")

        except Exception as e:
            logger.debug(f"Could not read pptx {path.name}: {e}")
            return ParsedDocument(
                file_path=str(path.resolve()),
                file_name=path.name,
                extension=path.suffix.lower(),
                file_size=stat.st_size,
                last_modified=stat.st_mtime,
                content_hash=file_hash,
                full_text=f"[Unreadable PowerPoint Document: {e}]",
                pages=[DocumentPage(page_number=1, text=f"[Unreadable: {e}]")],
                is_scanned=False,
                metadata={"error": str(e)}
            )

        full_text = "\n\n".join(full_text_parts)

        return ParsedDocument(
            file_path=str(path.resolve()),
            file_name=path.name,
            extension=path.suffix.lower(),
            file_size=stat.st_size,
            last_modified=stat.st_mtime,
            content_hash=file_hash,
            full_text=full_text,
            pages=pages,
            is_scanned=False,
            metadata={"slide_count": slide_count}
        )
